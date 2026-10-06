"""Pembuat dokumen Risalah Rapat (.md dan .docx).

Struktur mengikuti Risalah Rapat yang diekspor Sinopsis (bw-recorder
rapat-download.tsx), tetapi bagian Peserta dan Pokok Bahasan diisi dari hasil
diarization + notulensi terstruktur (Tabel 3.1 proposal), termasuk PIC dan
menit ke- sehingga setiap keputusan dapat ditelusuri ke rekamannya.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional

HARI = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
         "Agustus", "September", "Oktober", "November", "Desember"]


@dataclass
class RisalahStyle:
    kop: List[str] = field(default_factory=lambda: ["RISALAH RAPAT"])
    penanda_tangan: List[str] = field(default_factory=lambda: ["Dibuat Oleh", "Diperiksa Oleh", "Disetujui Oleh"])
    sertakan_transkrip: bool = True


@dataclass
class MeetingInfo:
    judul: str
    tanggal: Optional[str] = None  # YYYY-MM-DD
    tempat: str = ""
    mulai: Optional[str] = None  # HH:MM
    durasi_detik: float = 0.0


def format_tanggal(iso: Optional[str]) -> str:
    if not iso:
        return "-"
    d = date.fromisoformat(iso)
    return f"{HARI[d.weekday()]}, {d.day} {BULAN[d.month - 1]} {d.year}"


def format_pukul(mulai: Optional[str], durasi_detik: float) -> str:
    menit = round(durasi_detik / 60)
    if not mulai:
        return f"{menit} menit" if menit else "-"
    start = datetime.strptime(mulai, "%H:%M")
    end = datetime.fromtimestamp(start.timestamp() + durasi_detik)
    return f"{start:%H.%M} – {end:%H.%M} ({menit} menit)"


def gender_label(gender: Optional[str]) -> str:
    return {"pria": "Pria", "wanita": "Wanita"}.get(gender or "", "-")


def render_markdown(info: MeetingInfo, notulensi: dict, transcript: str = "", style: RisalahStyle = RisalahStyle()) -> str:
    out = [f"# {' · '.join(style.kop)}", ""]
    out += [
        f"**Agenda:** {info.judul}  ",
        f"**Hari, Tanggal:** {format_tanggal(info.tanggal)}  ",
        f"**Pukul:** {format_pukul(info.mulai, info.durasi_detik)}  ",
        f"**Tempat:** {info.tempat or '-'}",
        "",
        "## Peserta Rapat",
        "",
        "| No | Nama | Gender | Porsi bicara |",
        "|---|---|---|---|",
    ]
    for i, p in enumerate(notulensi.get("peserta", []), 1):
        out.append(f"| {i} | {p['nama']} | {gender_label(p.get('gender'))} | {p.get('porsi', 0):.0%} |")

    out += ["", "## Ringkasan Eksekutif", "", notulensi.get("ringkasan_eksekutif") or "_Belum tersedia._", ""]

    topik = notulensi.get("topik_pembahasan") or []
    out += ["## Topik Pembahasan", ""]
    out += [f"{i}. {t}" for i, t in enumerate(topik, 1)] or ["_Belum tersedia._"]

    out += ["", "## Pokok Bahasan, Keputusan & Tindak Lanjut", ""]
    points = notulensi.get("pokok_bahasan") or []
    if points:
        out += ["| No | Topik | Pembahasan | Keputusan | Tindak lanjut | PIC | Menit ke- |", "|---|---|---|---|---|---|---|"]
        for i, p in enumerate(points, 1):
            cells = [p.get(k, "") or "-" for k in ("topik", "pembahasan", "keputusan", "tindak_lanjut", "pic", "waktu")]
            out.append(f"| {i} | " + " | ".join(c.replace("|", "/") for c in cells) + " |")
    else:
        out.append(f"_{notulensi.get('catatan') or 'Belum tersedia.'}_")

    out += ["", "| " + " | ".join(style.penanda_tangan) + " |", "|" + "---|" * len(style.penanda_tangan)]
    out += ["| " + " | ".join(["<br><br>Nama:<br>Jabatan:"] * len(style.penanda_tangan)) + " |"]

    if style.sertakan_transkrip and transcript:
        out += ["", "---", "", "## Lampiran — Transkrip Berlabel Pembicara", "", "```", transcript, "```"]
    return "\n".join(out) + "\n"


def render_docx(info: MeetingInfo, notulensi: dict, path: Path, transcript: str = "", style: RisalahStyle = RisalahStyle()) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for i, line in enumerate(style.kop):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(16 if i == len(style.kop) - 1 else 12)

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    fields = [
        ("Agenda", info.judul),
        ("Hari, Tanggal", format_tanggal(info.tanggal)),
        ("Pukul", format_pukul(info.mulai, info.durasi_detik)),
        ("Tempat", info.tempat or "-"),
    ]
    for cell, (label, value) in zip([c for row in table.rows for c in row.cells], fields):
        p = cell.paragraphs[0]
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    doc.add_heading("Peserta Rapat", level=2)
    peserta = notulensi.get("peserta", [])
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, ["No", "Nama", "Gender", "Porsi bicara"]):
        cell.paragraphs[0].add_run(text).bold = True
    for i, p in enumerate(peserta, 1):
        row = t.add_row().cells
        for cell, text in zip(row, [str(i), p["nama"], gender_label(p.get("gender")), f"{p.get('porsi', 0):.0%}"]):
            cell.text = text

    doc.add_heading("Ringkasan Eksekutif", level=2)
    doc.add_paragraph(notulensi.get("ringkasan_eksekutif") or "Belum tersedia.")

    doc.add_heading("Topik Pembahasan", level=2)
    topik = notulensi.get("topik_pembahasan") or []
    for item in topik:
        doc.add_paragraph(item, style="List Number")
    if not topik:
        doc.add_paragraph("Belum tersedia.")

    doc.add_heading("Pokok Bahasan, Keputusan & Tindak Lanjut", level=2)
    points = notulensi.get("pokok_bahasan") or []
    if points:
        headers = ["No", "Topik", "Pembahasan", "Keputusan", "Tindak lanjut", "PIC", "Menit ke-"]
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for cell, text in zip(t.rows[0].cells, headers):
            cell.paragraphs[0].add_run(text).bold = True
        for i, p in enumerate(points, 1):
            values = [str(i)] + [p.get(k, "") or "-" for k in ("topik", "pembahasan", "keputusan", "tindak_lanjut", "pic", "waktu")]
            for cell, text in zip(t.add_row().cells, values):
                cell.text = text
    else:
        doc.add_paragraph(notulensi.get("catatan") or "Belum tersedia.")

    doc.add_paragraph()
    t = doc.add_table(rows=2, cols=len(style.penanda_tangan))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, style.penanda_tangan):
        cell.paragraphs[0].add_run(text).bold = True
    for cell in t.rows[1].cells:
        cell.text = "\n\n\nNama:\nJabatan:"

    if style.sertakan_transkrip and transcript:
        doc.add_page_break()
        doc.add_heading("Lampiran — Transkrip Berlabel Pembicara", level=2)
        for line in transcript.splitlines():
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(2)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
