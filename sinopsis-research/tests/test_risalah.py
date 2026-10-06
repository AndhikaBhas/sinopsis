from sinopsis_research.risalah import MeetingInfo, RisalahStyle, format_pukul, format_tanggal, render_docx, render_markdown

NOTULENSI = {
    "peserta": [
        {"label": "SPEAKER_01", "nama": "Sari Dewi", "gender": "wanita", "porsi": 0.55},
        {"label": "SPEAKER_00", "nama": "Budi Santoso", "gender": "pria", "porsi": 0.45},
    ],
    "topik_pembahasan": ["Pengadaan switch"],
    "pokok_bahasan": [
        {"topik": "Switch", "pembahasan": "Kebutuhan perangkat", "keputusan": "12 unit",
         "tindak_lanjut": "Kirim PO", "pic": "Sari Dewi", "waktu": "00:12:40"}
    ],
    "ringkasan_eksekutif": "Rapat menyetujui pengadaan 12 unit switch.",
}
INFO = MeetingInfo(judul="Evaluasi Pengadaan Q3", tanggal="2026-10-06", tempat="Ruang Rapat Lt. 3", mulai="09:02", durasi_detik=2700)


def test_format_tanggal_dan_pukul():
    assert format_tanggal("2026-10-06") == "Selasa, 6 Oktober 2026"
    assert format_pukul("09:02", 2700) == "09.02 – 09.47 (45 menit)"
    assert format_pukul(None, 2700) == "45 menit"


def test_markdown_contains_participants_points_and_transcript():
    md = render_markdown(INFO, NOTULENSI, transcript="[00:12:40] Sari Dewi: dua belas unit", style=RisalahStyle(kop=["PT Contoh", "RISALAH RAPAT"]))
    assert md.startswith("# PT Contoh · RISALAH RAPAT")
    assert "| 1 | Sari Dewi | Wanita | 55% |" in md
    assert "| Kirim PO | Sari Dewi | 00:12:40 |" in md
    assert "Lampiran" in md


def test_markdown_without_llm_shows_note():
    md = render_markdown(INFO, {"peserta": NOTULENSI["peserta"], "catatan": "Ollama belum terpasang"})
    assert "_Ollama belum terpasang_" in md


def test_docx_is_written(tmp_path):
    path = tmp_path / "risalah.docx"
    render_docx(INFO, NOTULENSI, path, transcript="[00:12:40] Sari Dewi: dua belas unit")
    from docx import Document

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Ringkasan Eksekutif" in text
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Sari Dewi" in cells and "00:12:40" in cells
